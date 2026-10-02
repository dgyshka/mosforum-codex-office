"""Exercise document tools on disposable files; never open the user's documents."""
import importlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile


def run(args, **kwargs):
    result = subprocess.run([str(a) for a in args], capture_output=True,
                            text=True, encoding='utf-8', errors='replace', timeout=180, **kwargs)
    if result.returncode:
        raise RuntimeError(f'{args[0]}: {result.stderr or result.stdout}')
    return result.stdout


def verify():
    for name in ['openpyxl', 'pandas', 'docx', 'pptx', 'pypdf', 'pdfplumber',
                 'pypdfium2', 'reportlab', 'pdf2image', 'pytesseract', 'PIL',
                 'numpy', 'defusedxml', 'lxml', 'markitdown']:
        importlib.import_module(name)
    for name in ['soffice', 'pandoc', 'pdftoppm', 'qpdf', 'tesseract', 'node', 'git', 'jq']:
        if not shutil.which(name):
            raise RuntimeError(f'Программа не найдена: {name}')
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from pptx import Presentation
    from openpyxl import Workbook, load_workbook
    from pypdf import PdfReader
    from markitdown import MarkItDown
    from PIL import Image, ImageDraw, ImageFont

    with tempfile.TemporaryDirectory(prefix='office-check-', ignore_cleanup_errors=True) as folder:
        work = Path(folder)
        doc = Document()
        doc.add_paragraph('Проверка документов 42')
        doc.save(work/'word.docx')
        assert 'Проверка' in run(['pandoc', work/'word.docx', '-t', 'plain'])
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = 'Проверка презентации 42'
        deck.save(work/'slides.pptx')
        for filename in ['word.docx', 'slides.pptx']:
            run(['soffice', '--headless', f'-env:UserInstallation={(work/"lo").as_uri()}',
                 '--convert-to', 'pdf', '--outdir', work, work/filename])
            pdf = work/Path(filename).with_suffix('.pdf')
            assert '42' in ''.join(p.extract_text() for p in PdfReader(pdf).pages), filename
            run(['qpdf', '--check', pdf])
        run(['pdftoppm', '-f', '1', '-singlefile', '-png', work/'word.pdf', work/'page'])
        assert (work/'page.png').stat().st_size > 0
        print('Word, презентации, PDF и просмотр страниц: проверено.')

        tracked = Document()
        paragraph = tracked.add_paragraph('Проверка ')
        insertion = OxmlElement('w:ins')
        insertion.set(qn('w:id'), '1')
        insertion.set(qn('w:author'), 'Office check')
        insertion.set(qn('w:date'), '2026-01-01T00:00:00Z')
        text_run, text_node = OxmlElement('w:r'), OxmlElement('w:t')
        text_node.text = '42'
        text_run.append(text_node)
        insertion.append(text_run)
        paragraph._p.append(insertion)
        tracked.save(work/'tracked.docx')
        run([sys.executable, Path(__file__).parent/'profile/skills/docx/scripts/accept_changes.py',
             work/'tracked.docx', work/'accepted.docx'])
        accepted = Document(work/'accepted.docx')
        assert not accepted._element.xpath('//w:ins')
        assert '42' in ''.join(p.text for p in accepted.paragraphs)
        assert Document(work/'tracked.docx')._element.xpath('//w:ins')
        print('Word: принятие исправлений в копии документа проверено.')

        book = Workbook()
        book.active['A1'] = 21
        book.active['A2'] = '=A1*2'
        book.save(work/'table.xlsx')
        script = Path(__file__).parent/'profile/skills/xlsx/scripts/recalc.py'
        result = json.loads(run([sys.executable, script, work/'table.xlsx', '120']))
        assert result.get('status') == 'success', result
        cached = load_workbook(work/'table.xlsx', data_only=True)
        assert cached.active['A2'].value == 42
        cached.close()
        assert '42' in MarkItDown().convert(str(work/'table.xlsx')).text_content
        print('Excel: формула пересчитана, результат 42 прочитан из файла.')

        assert {'rus', 'eng', 'osd'}.issubset(run(['tesseract', '--list-langs']).split())
        font = ImageFont.truetype(str(Path(os.environ.get('WINDIR', 'C:/Windows'))/'Fonts/arial.ttf'), 64)
        picture = Image.new('RGB', (850, 150), 'white')
        ImageDraw.Draw(picture).text((20, 25), 'Проверка 42', font=font, fill='black')
        picture.save(work/'scan.png')
        recognized = run(['tesseract', work/'scan.png', 'stdout', '-l', 'rus', '--psm', '7'])
        assert '42' in recognized and 'проверка' in recognized.lower(), recognized
        english = run(['tesseract', work/'scan.png', 'stdout', '-l', 'eng', '--psm', '7'])
        assert '42' in english, english
        print('Распознавание русского текста, английский язык и osd: проверено.')

        js = """const fs=require('fs');
const {Document,Paragraph,Packer}=require('docx');
const PptxGenJS=require('pptxgenjs');
require('react-icons'); require('react'); require('react-dom'); require('sharp');
(async()=>{
fs.writeFileSync('node.docx',await Packer.toBuffer(new Document({sections:[{children:[new Paragraph('42')]}]})));
const p=new PptxGenJS();p.addSlide().addText('42',{x:1,y:1,w:3,h:1});await p.writeFile({fileName:'node.pptx'});
})().catch(e=>{console.error(e);process.exit(1)});
"""
        (work/'check.cjs').write_text(js, encoding='utf-8')
        run(['node', work/'check.cjs'], cwd=work)
        assert Document(work/'node.docx').paragraphs[0].text == '42'
        assert len(Presentation(work/'node.pptx').slides) == 1
        print('Библиотеки Python и Node: проверено создание документов.')


if __name__ == '__main__':
    try:
        verify()
    except Exception as error:
        print('Проверка офисных инструментов не пройдена: ' + str(error), file=sys.stderr)
        sys.exit(1)
